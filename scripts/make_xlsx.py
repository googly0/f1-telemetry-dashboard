"""Race-by-race backtest workbook (called by backtest.py --xlsx). Summary tabs are live formulas."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L

import f1model as m

ROOT = Path(__file__).resolve().parent.parent
FONT = "Arial"
H = Font(name=FONT, bold=True, color="FFFFFF")
HF = PatternFill("solid", fgColor="1F3A5F")
B = Font(name=FONT)
BOLD = Font(name=FONT, bold=True)
BLUE = Font(name=FONT, color="0000FF")
KEY = PatternFill("solid", fgColor="FFF2CC")
NOTE = Font(name=FONT, italic=True, size=9, color="555555")

MODELS = [("C", "Race-day"), ("F", "Pre-quali"), ("A", "Grid only")]
LABELS = {
    "x_grid": "Starting grid (-ln position)", "x_pole": "Pole position (0/1)",
    "x_qgap": "Qualifying gap to pole (-%)", "x_sprint": "Sprint result (-ln position)",
    "x_team_form": "Team points, last 6 races (1 = winning every race)",
    "x_team_qgap": "Team qualifying gap, last 6 races (-%)",
    "x_team_gain": "Team places gained on race day, last 6 (/5)",
    "x_tm": "Finishing places ahead of teammate, last 6 (/5)",
    "x_elo_d": "Driver Elo (/100)", "x_elo_t": "Team Elo (/100)",
}


def _hdr(ws, row, vals, height=None):
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font, c.fill = H, HF
        c.alignment = Alignment(wrap_text=True, vertical="center")
    if height:
        ws.row_dimensions[row].height = height


def write(path, F, board, per_race):
    years = json.loads((ROOT / "lib/data/all_years.json").read_text())["years"]
    names = {}
    race_names = {}
    for y, yd in years.items():
        for k, v in yd["names"]["drivers"].items():
            names[int(k)] = v
        for r in yd["rounds"]:
            race_names[(int(y), r["round"])] = r["name"]
    grid = {(r.raceId, r.driverId): r.grid for r in F[["raceId", "driverId", "grid"]].itertuples(index=False)}

    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    # ------------------------------------------------------------------ Races
    wr = wb.active
    wr.title = "Races"
    cols = ["Year", "Round", "Race", "Split", "Winner", "Winner's grid slot"]
    for _, lab in MODELS:
        cols += [f"{lab}: pick", f"{lab}: P(pick)", f"{lab}: 2nd pick", f"{lab}: 3rd pick", f"{lab}: P(actual winner)",
                 f"{lab}: hit", f"{lab}: winner in top 3", f"{lab}: podium correct (0-3)", f"{lab}: log-loss"]
    _hdr(wr, 1, cols, 48)
    frames = {}
    for key, _ in MODELS:
        t = per_race[key]["train"].assign(split="Train")
        s = per_race[key]["test"].assign(split="Test")
        frames[key] = __import__("pandas").concat([t, s]).set_index("raceId")
    base = frames["C"].sort_values(["year", "round"])
    row = 2
    for rid, r in base.iterrows():
        vals = [int(r.year), int(r["round"]), race_names.get((int(r.year), int(r["round"])), ""), r.split,
                names.get(int(r.winner), str(r.winner)), int(grid.get((rid, int(r.winner)), 0))]
        for j, (key, _) in enumerate(MODELS):
            x = frames[key].loc[rid]
            c0 = 7 + j * 9  # first column of this model's block
            pick, p2, p3 = (L(c0), L(c0 + 2), L(c0 + 3))
            pod = len(set(x.podium_pick.split(",")) & set(x.podium_actual.split(",")))
            vals += [names.get(int(x.pick), str(x.pick)), float(x.p_pick),
                     names.get(int(x.pick2), str(x.pick2)), names.get(int(x.pick3), str(x.pick3)),
                     float(x.p_winner),
                     f"=IF({pick}{row}=$E{row},1,0)",
                     f"=IF(OR({pick}{row}=$E{row},{p2}{row}=$E{row},{p3}{row}=$E{row}),1,0)",
                     pod, f"=-LN(MAX({L(c0 + 4)}{row},0.0001))"]
        for i, v in enumerate(vals, 1):
            c = wr.cell(row=row, column=i, value=v)
            c.font = B
            if cols[i - 1].endswith(("P(pick)", "P(actual winner)")):
                c.number_format = "0.0%"
            if cols[i - 1].endswith("log-loss"):
                c.number_format = "0.000"
        row += 1
    last = row - 1
    wr.freeze_panes = "E2"
    for i, w in enumerate([6, 6, 26, 7, 18, 8], 1):
        wr.column_dimensions[L(i)].width = w
    for i in range(7, len(cols) + 1):
        wr.column_dimensions[L(i)].width = 11 if "pick" not in cols[i - 1] or "P(" in cols[i - 1] else 17
    wr["D1"].comment = Comment("Train = 2016-2023, the seasons the weights were fitted on (in-sample). "
                               "Test = 2024 onward, never seen during fitting.", "backtest")
    wr["F1"].comment = Comment("Starting grid of the eventual winner (qualifying position where the source has no grid).", "backtest")

    def rng(col):
        return f"Races!${col}$2:${col}${last}"

    # ------------------------------------------------------------------ Summary
    ws = wb.create_sheet("Summary", 0)
    ws["A1"] = "F1 Predictor — Held-out Backtest"
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = ("Weights fitted on 2016–2023, then every race from 2024 to the latest round predicted blind. "
                "Each prediction uses only what was known before lights out. All numbers below are formulas over the Races tab.")
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A2:E2")
    ws.row_dimensions[2].height = 42
    metric_rows = [("Favourite actually wins", "hit", "0.0%"), ("Winner in model's top 3", "top3", "0.0%"),
                   ("Podium drivers predicted", "podium", "0.0%"), ("Win log-loss (lower = better)", "ll", "0.000")]

    def block(top, split, title):
        _hdr(ws, top, [title, "Race-day", "Pre-quali", "Grid only", "How to read it"], 30)
        ws.cell(row=top + 1, column=1, value="Races").font = B
        for j in range(3):
            c = ws.cell(row=top + 1, column=2 + j, value=f'=COUNTIFS({rng("D")},"{split}")')
            c.font = B
        explain = {"hit": "The model's most likely winner won", "top3": "Winner was one of the model's 3 most likely winners",
                   "podium": "Share of the dashboard's predicted podium that finished on the podium",
                   "ll": "Average -ln P(actual winner); uniform guessing ≈ 3.1"}
        for i, (lab, key, fmt) in enumerate(metric_rows):
            r = top + 2 + i
            ws.cell(row=r, column=1, value=lab).font = BOLD if key == "top3" else B
            for j in range(3):
                c0 = 7 + j * 9
                col = {"hit": L(c0 + 5), "top3": L(c0 + 6), "podium": L(c0 + 7), "ll": L(c0 + 8)}[key]
                f = f'=AVERAGEIFS({rng(col)},{rng("D")},"{split}")' + ("/3" if key == "podium" else "")
                c = ws.cell(row=r, column=2 + j, value=f)
                c.font, c.number_format = B, fmt
                if key == "top3":
                    c.fill = KEY
            ws.cell(row=r, column=5, value=explain[key]).font = NOTE
        return top + 2 + len(metric_rows)

    end = block(4, "Test", "Held-out test (2024 onward)")
    end = block(end + 1, "Train", "Training seasons 2016–2023 (in-sample)")
    r = end + 1
    ws.cell(row=r, column=1, value="Résumé line this supports").font = BOLD
    ws.cell(row=r + 1, column=1, value=(
        '="Plackett-Luce race model backtested on "&B5&" held-out races (2024–2026): winner in top-3 picks "'
        '&TEXT(B7,"0%")&", favourite won "&TEXT(B6,"0%")&", win log-loss "&TEXT(B9,"0.00")&" vs 3.1 random."')).font = B
    ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=5)
    ws.cell(row=r + 1, column=1).alignment = Alignment(wrap_text=True)
    ws.row_dimensions[r + 1].height = 32
    ws.cell(row=r + 3, column=1, value=(
        "Yellow row: the honest '85%+' number. Calling the exact winner tops out near 55–60% for any pre-race model; "
        "the starting grid alone gets 58.7% on these seasons.")).font = NOTE
    ws.merge_cells(start_row=r + 3, start_column=1, end_row=r + 3, end_column=5)
    ws.cell(row=r + 3, column=1).alignment = Alignment(wrap_text=True)
    ws.row_dimensions[r + 3].height = 30
    for i, w in enumerate([34, 13, 13, 13, 52], 1):
        ws.column_dimensions[L(i)].width = w

    # ------------------------------------------------------------------ Arena
    wa = wb.create_sheet("Arena", 1)
    wa["A1"] = "Model arena — every candidate, including the ones that lost"
    wa["A1"].font = Font(name=FONT, bold=True, size=13)
    _hdr(wa, 3, ["", "Model", "Val: favourite wins", "Val: winner in top 3", "Val: log-loss",
                 "Test: favourite wins", "Test: winner in top 3", "Test: podium", "Test: log-loss"], 32)
    for i, b in enumerate(board.itertuples(index=False), 4):
        vals = [b.key, b.model, b.val_hit, b.val_top3, b.val_ll, b.test_hit, b.test_top3, b.test_podium, b.test_ll]
        for j, v in enumerate(vals, 1):
            c = wa.cell(row=i, column=j, value=round(float(v), 4) if isinstance(v, float) else v)
            c.font = BLUE if j > 2 else B
            c.number_format = "0.000" if j in (5, 9) else "0.0%"
    n = 4 + len(board)
    wa.cell(row=n + 1, column=1, value=(
        "Blue = values computed by scripts/backtest.py (not formulas). Validation = rolling seasons 2018–2023, each fit only "
        "on earlier seasons; it was used to choose features. Test = fit 2016–2023, predict 2024+, reported once. Model D won "
        "validation but lost the test: Elo history learned to trust dominant cars (Mercedes, then Red Bull) and broke when "
        "McLaren rose in 2024 and the 2026 rules reset the order, so the shipped race-day model (C) uses this-season form only.")).font = NOTE
    wa.merge_cells(start_row=n + 1, start_column=1, end_row=n + 1, end_column=9)
    wa.cell(row=n + 1, column=1).alignment = Alignment(wrap_text=True)
    wa.row_dimensions[n + 1].height = 70
    for i, w in enumerate([4, 44, 12, 12, 10, 12, 12, 10, 10], 1):
        wa.column_dimensions[L(i)].width = w

    # ------------------------------------------------------------------ By Year
    wy = wb.create_sheet("By Year", 2)
    hdr = ["Season", "Races"]
    for _, lab in MODELS:
        hdr += [f"{lab}: favourite wins", f"{lab}: winner in top 3", f"{lab}: log-loss"]
    _hdr(wy, 1, hdr, 32)
    yrs = sorted(base.year.unique())
    for i, y in enumerate(yrs, 2):
        wy.cell(row=i, column=1, value=str(int(y))).font = B
        c = wy.cell(row=i, column=2, value=f'=COUNTIFS({rng("A")},VALUE(A{i}))')
        c.font = B
        for j in range(3):
            c0 = 7 + j * 9
            for k, col in enumerate((L(c0 + 5), L(c0 + 6), L(c0 + 8))):
                cell = wy.cell(row=i, column=3 + j * 3 + k, value=f'=AVERAGEIFS({rng(col)},{rng("A")},VALUE(A{i}))')
                cell.font, cell.number_format = B, ("0.000" if k == 2 else "0.0%")
    wy.cell(row=len(yrs) + 3, column=1, value="2016–2023 are in-sample (training seasons); 2024 onward is held out.").font = NOTE
    for i in range(1, len(hdr) + 1):
        wy.column_dimensions[L(i)].width = 12
    wy.freeze_panes = "B2"

    # ------------------------------------------------------------------ Model
    wm = wb.create_sheet("Model", 3)
    wm["A1"] = "Fitted weights (Plackett-Luce, maximum likelihood on 2016–2023)"
    wm["A1"].font = Font(name=FONT, bold=True, size=13)
    r = 3
    for key, title, fit_note in (("C", "Race-day model", "fit on race winners (k=1)"),
                                 ("F", "Pre-qualifying model", "fit on the whole classified order, position j weighted 0.3^j")):
        _hdr(wm, r, [f"{title} — {fit_note}", "Weight"], 30)
        for i, (f, w) in enumerate(zip(per_race[key]["cols"], per_race[key]["beta"]), r + 1):
            wm.cell(row=i, column=1, value=LABELS.get(f, f)).font = B
            c = wm.cell(row=i, column=2, value=round(float(w), 4))
            c.font, c.number_format = BLUE, "0.000"
        r += len(per_race[key]["cols"]) + 2
    wm.cell(row=r, column=1, value=(
        "Strength = sum of weight × feature. A driver's chance of beating the rest of the remaining field is proportional to "
        "exp(strength). Every car first retires with its team's recent DNF rate (last ~30 starts, shrunk toward 12%); "
        "the finishers are then ordered by Plackett-Luce.")).font = NOTE
    wm.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    wm.cell(row=r, column=1).alignment = Alignment(wrap_text=True)
    wm.row_dimensions[r].height = 60
    wm.column_dimensions["A"].width = 60
    wm.column_dimensions["B"].width = 12
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
